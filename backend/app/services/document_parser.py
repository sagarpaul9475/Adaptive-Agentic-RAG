from pathlib import Path
from docx import Document as DocxDocument
from pypdf import PdfReader

class ParsedDocument:
    def __init__(self, text: str, page_count: int | None = None):
        self.text = text
        self.page_count = page_count

def parse_document(file_path: Path, extension: str) -> ParsedDocument:
    extension = extension.lower().lstrip(".")
    if extension == "pdf":
        reader = PdfReader(str(file_path))
        pages = []
        for page in reader.pages:
            page_text = page.extract_text() or ""
            if page_text.strip():
                pages.append(page_text)
        return ParsedDocument("\n\n".join(pages), len(reader.pages))
    if extension in {"txt", "md"}:
        return ParsedDocument(file_path.read_text(encoding="utf-8", errors="replace"))
    if extension == "docx":
        document = DocxDocument(str(file_path))
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        return ParsedDocument("\n\n".join(paragraphs))
    raise ValueError(f"Unsupported document type: .{extension}")
